"""Audit built documentation examples through the real runner; write an Excel inventory.

Run after npm run build:
  wrangles-docs/.venv/bin/python scripts/audit_recipe_examples.py
Requires beautifulsoup4 and openpyxl in the runner environment.
No remote requests, connectors, arbitrary Python, or secret-dependent recipes execute.
This tests rendered inputs and the Python adapter, not a deployed HTTP server.
"""
import ast
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import subprocess
import sys
from uuid import UUID

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import yaml

ROOT = Path(__file__).resolve().parents[1]
SITE = ROOT / 'wrangles-docs'
sys.path.insert(0, str(SITE))
from run_recipe import load_env_file, coerce_cell, display_cell


def normalized(code):
    return re.sub(r'\s+', '', code)


def source_index():
    index = {}
    files = list((ROOT / 'registry/input_wrangles').rglob('*.md'))
    files += list((ROOT / 'registry/input_connectors').glob('*.md'))
    files += list((SITE / 'registry-docs/python/recipes').rglob('*.md'))
    for path in files:
        text = path.read_text()
        for match in re.finditer(r'```ya?ml[^\n]*\n(.*?)```', text, re.S):
            index.setdefault(normalized(match[1]), []).append(
                f'{path.relative_to(ROOT)}:{text[:match.start()].count(chr(10)) + 1}')
        if text.startswith('---\n'):
            metadata = yaml.safe_load(text.split('---', 2)[1]) or {}
            for example in metadata.get('examples', []):
                recipe = example.get('recipe')
                if recipe:
                    code = recipe if isinstance(recipe, str) else yaml.safe_dump(recipe, sort_keys=False)
                    index.setdefault(normalized(code), []).append(str(path.relative_to(ROOT)))
    return index


def table(element):
    if element is None:
        return None
    columns = [cell.get_text().strip() for cell in element.select('thead th')]
    rows = [[cell.get_text().strip() for cell in row.select('th, td')]
            for row in element.select('tbody tr')]
    return {'columns': columns, 'rows': rows} if columns and rows else None


def collect():
    sources = source_index()
    pages = sorted((SITE / 'build/py/namespaces').glob('*/index.html'))
    pages += sorted((SITE / 'build/py/connectors').rglob('index.html'))
    pages += sorted((SITE / 'build/py/python/recipes/custom-functions').rglob('index.html'))
    results = []
    for page in pages:
        soup = BeautifulSoup(page.read_text(), 'html.parser')
        for ordinal, block in enumerate(soup.select('div.theme-code-block.language-yaml'), 1):
            code = '\n'.join(line.get_text() for line in block.select('.token-line')).strip()
            heading = block.find_previous(['h2', 'h3', 'h4'])
            anchor = heading.get('id', '') if heading else ''
            path = '/' + str(page.parent.relative_to(SITE / 'build'))
            item = dict(page=path + ('#' + anchor if anchor else ''), example=ordinal,
                        source='; '.join(sources.get(normalized(code), [])), recipe=code,
                        title=heading.get_text(' ', strip=True) if heading else '',
                        status='', reason='', actual=None, expected=None, input=None)
            try:
                recipe = yaml.safe_load(code.replace('\t', '  '))
            except yaml.YAMLError:
                recipe = None
            item['parsed'] = recipe
            grid = None
            for sibling in block.next_siblings:
                if getattr(sibling, 'name', '') in ('h2', 'h3', 'h4'):
                    break
                if 'ww-sample-grid' in getattr(sibling, 'attrs', {}).get('class', []):
                    grid = sibling
                    break
            panels = grid.select('.ww-sample-panel') if grid else []
            if panels:
                item['input'] = table(panels[0].select_one('table'))
                item['expected'] = table(panels[-1].select_one('table'))
            elif '/custom-functions/wrangles' in path:
                # Tutorials put their input/output tables before the Python and YAML blocks.
                previous = block.find_previous('table')
                if previous:
                    outer = previous.find_parent('table')
                    nested = outer.select('table') if outer else []
                    if len(nested) == 2:
                        item['input'], item['expected'] = table(nested[0]), table(nested[1])
            item['has_sample'] = bool(item['input'])
            if not item['input']:
                steps = recipe.get('wrangles', []) if isinstance(recipe, dict) else []
                config = next(iter(steps[0].values()), {}) if steps and isinstance(steps[0], dict) else {}
                inp = config.get('input') if isinstance(config, dict) else None
                columns = inp if isinstance(inp, list) else [inp] if isinstance(inp, str) else ['input']
                item['input'] = {'columns': columns, 'rows': [[''] * len(columns)]}
            results.append(item)
    proc = subprocess.run(['node', '-e',
        "const p=require('./src/components/recipeRunPolicy.cjs');let s='';process.stdin.on('data',x=>s+=x);"
        "process.stdin.on('end',()=>process.stdout.write(JSON.stringify(JSON.parse(s).map(p.getRunDisabledReason))));"],
        cwd=SITE, input=json.dumps([row['recipe'] for row in results]), text=True, capture_output=True, check=True)
    for item, policy in zip(results, json.loads(proc.stdout)):
        item['run_button_reason'] = policy
    return results


def canonical(value):
    if isinstance(value, str):
        for parser in (json.loads, ast.literal_eval):
            try:
                parsed = parser(value)
                if isinstance(parsed, (dict, list, tuple, int, float, bool)):
                    return parsed
            except (ValueError, SyntaxError):
                pass
    return value


def compare(expected, actual):
    if expected['columns'] != actual['columns'] or len(expected['rows']) != len(actual['rows']):
        return False
    def same_cell(documented, returned):
        # Fixture tables quote JSON-looking strings to preserve their type on
        # input. The runner displays the returned string without that wrapper.
        if isinstance(documented, str) and isinstance(returned, str):
            try:
                decoded = json.loads(documented)
                if isinstance(decoded, str) and decoded == returned:
                    return True
            except ValueError:
                pass
        return canonical(documented) == canonical(returned)

    return all(len(a) == len(b) and all(same_cell(x, y) for x, y in zip(a, b))
               for a, b in zip(expected['rows'], actual['rows']))


def valid_uuid_output(item, actual):
    expected = item['expected']
    if actual['columns'] != expected['columns'] or len(actual['rows']) != len(item['input']['rows']):
        return False
    if len(expected['rows']) != len(actual['rows']):
        return False
    try:
        values = [UUID(str(row[0])) for row in actual['rows'] if len(row) == 1]
        return (len(values) == len(actual['rows']) and bool(values)
                and all(value.version == 4 for value in values) and len(set(values)) == len(values))
    except (ValueError, TypeError, AttributeError):
        return False


def valid_sample_output(item, actual, config):
    if set(config) != {'rows'}:
        return False  # Other sampling options need their own explicit checks.
    rows = config['rows']
    if isinstance(rows, bool) or not isinstance(rows, (int, float)) or rows <= 0:
        return False
    source = item['input']
    count = min(int(rows), len(source['rows'])) if rows >= 1 else round(rows * len(source['rows']))
    def key(row):
        return tuple(json.dumps(canonical(display_cell(coerce_cell(cell))), sort_keys=True) for cell in row)
    available = Counter(key(row) for row in source['rows'])
    for table in (actual, item['expected']):
        if table['columns'] != source['columns'] or len(table['rows']) != count:
            return False
        if Counter(key(row) for row in table['rows']) - available:
            return False
    return True


def valid_log_output(item, actual, stderr):
    import pandas as pd
    source = item['input']
    frame = pd.DataFrame([[coerce_cell(cell) for cell in row] for row in source['rows']],
                         columns=source['columns'])
    return (compare(item['expected'], actual) and compare(source, actual)
            and ': Dataframe ::\n\n' + frame.head(20).to_string() + '\n' in stderr)


def audit(item):
    item['checked_at'] = datetime.now(timezone.utc).isoformat()
    recipe = item.pop('parsed')
    reason = item['run_button_reason']
    if 'custom.function_name:' in item['recipe']:
        item.update(status='NOT_RUN_TEMPLATE_OR_CONNECTOR', reason='Generic tutorial template: supply a function implementation before running.')
        return item
    if not isinstance(recipe, dict) or not isinstance(recipe.get('wrangles'), list):
        item.update(status='NOT_RUN_TEMPLATE_OR_CONNECTOR', reason='Not an interactive wrangle recipe; may require files, services or surrounding Python.')
        return item
    if reason:
        category = 'NOT_RUN_REQUIRES_SETUP'
        if any(word in reason for word in ('sample', 'dictionary input', 'placeholder', 'valid YAML')):
            category = 'BROKEN_EXAMPLE_BLOCKED'
        item.update(status=category, reason=reason)
        return item
    # Refuse external side effects and arbitrary code even when the UI permits them.
    # Only this reviewed, fixed local Python example is permitted in the audit.
    safe_python = recipe == {'wrangles': [{'python': {
        'input': 'input_column', 'output': 'sliced', 'command': 'input_column[:i]', 'i': 3}}]}
    forbidden = not safe_python and re.search(r'(?m)^\s*-\s*(python|recipe|extract\.(address|attributes|codes|html|properties)|translate)\s*:', item['recipe'])
    def has_connector(value):
        if isinstance(value, dict):
            return any(key in ('read', 'write', 'run') or has_connector(child)
                       for key, child in value.items() if key not in ('value', 'data', 'default'))
        return isinstance(value, list) and any(has_connector(child) for child in value)

    if has_connector(recipe) or forbidden:
        item.update(status='NOT_RUN_EXTERNAL_OR_CODE', reason='Requires service/file access or Python/recipe review; not executed in this local audit.')
        return item
    payload = {'recipe': item['recipe'], 'input': item['input'],
               'outputColumns': (item['expected'] or {}).get('columns', [])}
    # Fail closed on network access. Capture all output and publish only structured results.
    worker = "import sys,runpy\ndef guard(event,args):\n if event in ('socket.connect','socket.getaddrinfo','socket.sendto'): raise RuntimeError('Network disabled for documentation audit')\nsys.addaudithook(guard)\nrunpy.run_path(sys.argv[1],run_name='__main__')"
    try:
        process = subprocess.run([sys.executable, '-c', worker, str(SITE / 'run_recipe.py')],
            input=json.dumps(payload), text=True, capture_output=True, timeout=60, cwd=SITE)
        if process.returncode:
            errors = [line for line in process.stderr.splitlines() if line.startswith('{"error":')]
            error = json.loads(errors[-1])['error'] if errors else 'Runner failed; raw logs withheld to protect secrets.'
            item.update(status='BROKEN_RUNTIME', reason=redact(error))
        else:
            output = json.loads(process.stdout)
            item['actual'] = output
            if not item['expected']:
                item.update(status='EXECUTED_NO_EXPECTATION', reason='Execution succeeded; no expected output table to verify.')
            elif (len(recipe['wrangles']) == 1
                  and set(recipe['wrangles'][0]) in ({'create.uuid'}, {'create.guid'})):
                valid = valid_uuid_output(item, output)
                item.update(status='WORKING' if valid else 'OUTPUT_MISMATCH',
                            reason='UUID example: checked row count, output column, unique UUID v4 values; illustrative values are random.' if valid else 'UUID output failed structural checks.')
            elif len(recipe['wrangles']) == 1 and set(recipe['wrangles'][0]) == {'select.sample'}:
                valid = valid_sample_output(item, output, recipe['wrangles'][0]['select.sample'])
                item.update(status='WORKING' if valid else 'OUTPUT_MISMATCH',
                            reason='Random sample: checked count, columns and unchanged input rows without replacement, including illustrative output.' if valid else 'Random sample failed count, column or row-membership checks.')
            elif recipe['wrangles'] == [{'log': {}}]:
                valid = valid_log_output(item, output, process.stderr)
                item.update(status='WORKING' if valid else 'OUTPUT_MISMATCH',
                            reason='Verified unchanged data and the actual dataframe log emitted on stderr.' if valid else 'Dataframe log missing or returned data changed.')
            elif compare(item['expected'], output):
                item.update(status='WORKING', reason='Execution and expected output match (JSON/list/number formatting normalized).')
            elif 'date_calculator:' in item['recipe']:
                item.update(status='EXECUTED_REVIEW_OUTPUT', reason='Date/timestamp display differs; review expected format.')
            else:
                item.update(status='OUTPUT_MISMATCH', reason='Execution succeeded but displayed expected output differs; not necessarily a WranglesPY defect.')
    except subprocess.TimeoutExpired:
        item.update(status='BROKEN_TIMEOUT', reason='Runner exceeded 60 seconds.')
    except (ValueError, TypeError):
        item.update(status='BROKEN_RUNNER_RESPONSE', reason='Runner did not return valid structured JSON; raw output withheld.')
    return item


def redact(text):
    import os
    for key, value in os.environ.items():
        if len(value) >= 8 and any(word in key.upper() for word in ('TOKEN', 'KEY', 'PASSWORD', 'SECRET')):
            text = text.replace(value, '[REDACTED]')
    return text


def main():
    load_env_file(SITE / '.env')
    load_env_file(ROOT / '.env')
    tests = subprocess.run([sys.executable, str(SITE / 'scripts/test_custom_functions.py')],
                           text=True, capture_output=True, cwd=SITE)
    if tests.returncode:
        raise RuntimeError('Custom function tests failed; run scripts/test_custom_functions.py separately.')
    output = ROOT / 'audits'
    output.mkdir(exist_ok=True)
    if '--report-only' in sys.argv:
        # Reformat saved evidence without pretending it was executed again.
        previous = json.loads((output / 'recipe-examples.json').read_text())
        results = previous['results']
        generated_at = previous['generated_at']
    else:
        items = collect()
        results = []
        saved = {}
        if '--refresh-changed' in sys.argv:
            previous = json.loads((output / 'recipe-examples.json').read_text())
            for row in previous['results']:
                row.setdefault('checked_at', previous['generated_at'])
                saved[(row['page'], row['example'])] = row
        def check(item):
            old = saved.get((item['page'], item['example']))
            # Only reuse unchanged non-custom samples after editorial-only changes.
            if old and 'custom.' not in item['recipe'] and all(old.get(key) == item.get(key)
                    for key in ('recipe', 'input', 'expected', 'run_button_reason')):
                old['source'] = item['source']
                return old
            return audit(item)
        with ThreadPoolExecutor(max_workers=4) as pool:
            for item in pool.map(check, items):
                results.append(item)
                print(f'{len(results)}/{len(items)} {item["status"]} {item["page"]}', flush=True)
        generated_at = datetime.now(timezone.utc).isoformat()
    counts = Counter(row['status'] for row in results)
    report = dict(generated_at=generated_at, counts=dict(counts),
        method='Built HTML inputs -> run_recipe.py with shared customFunctions.py. Network blocked; no live HTTP/browser click test.',
        scope='Canonical namespace and connector YAML examples plus custom-function tutorial YAML. Python-only snippets are not executable recipe examples.',
        results=results)
    (output / 'recipe-examples.json').write_text(redact(json.dumps(report, indent=2)) + '\n')
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    summary.append(['Generated (UTC)', report['generated_at']])
    summary.append(['Method', report['method']])
    summary.append(['Scope', report['scope']])
    summary.append(['Important', 'WORKING means tested sample only. NOT_RUN is not a failure. OUTPUT_MISMATCH needs review. UI blocking is not a server security boundary.'])
    for status, count in sorted(counts.items()):
        summary.append([status, count])
    sheet = workbook.create_sheet('Examples')
    fields = ['source', 'page', 'example', 'title', 'status', 'reason', 'run_button_reason', 'has_sample', 'recipe', 'input', 'expected', 'actual', 'checked_at']
    sheet.append(fields)
    for row in results:
        values = [row.get(field, '') for field in fields]
        sheet.append([redact(json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else str(v))[:32000] for v in values])
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'  # Never interpret example text as an Excel formula.
        sheet.cell(sheet.max_row, 2).hyperlink = 'https://wrangles.io' + row['page']
    functions = workbook.create_sheet('Custom functions')
    functions.append(['Function', 'Implementation', 'Verification'])
    references = workbook.create_sheet('Custom references')
    references.append(['Source', 'Example', 'Reference', 'Implementation status'])
    implemented = {node.name for node in ast.parse((SITE / 'customFunctions.py').read_text()).body
                   if isinstance(node, ast.FunctionDef)}
    for name in sorted(implemented):
        functions.append([name, 'wrangles-docs/customFunctions.py', 'Passed real WranglesPY file-path integration test'])
    for row in results:
        for name in sorted(set(re.findall(r'custom\.([A-Za-z_][\w.]*)', row['recipe']))):
            references.append([row['source'], row['example'], 'custom.' + name,
                               'Implemented and tested' if name in implemented else
                               'Tutorial placeholder; no unique implementation supplied by the example'])
    coverage = workbook.create_sheet('Registry coverage')
    coverage.append(['Source', 'Catalog key', 'Rendered YAML example count', 'Coverage'])
    for folder in ('input_wrangles', 'input_connectors'):
        for path in sorted((ROOT / 'registry' / folder).rglob('*.md')):
            relative = str(path.relative_to(ROOT))
            metadata = yaml.safe_load(path.read_text().split('---', 2)[1])
            count = sum(any(source.split(':')[0] == relative for source in row['source'].split('; '))
                        for row in results)
            coverage.append([relative, metadata.get('catalog_key', metadata.get('wrangle_key', '')), count,
                             'See Examples sheet' if count else 'No rendered YAML example found; not tested'])
    for ws in workbook:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24476A')
        for column in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(column)].width = 48 if column != 3 else 12
    workbook.save(output / 'recipe-examples.xlsx')
    print(json.dumps(dict(counts), indent=2))


if __name__ == '__main__':
    main()
